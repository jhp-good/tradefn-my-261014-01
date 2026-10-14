"""1일차 결제조건 매핑표 — payment_terms_raw 표기 → 표준 코드 7개.

실행: uv run python scripts/d1_make_code_table.py
결과: workbench/day1/outputs/d1_code_table.csv (utf-8-sig, 엑셀에서 한글이 깨지지 않게)
원본(data/)은 읽기만 한다. 판단은 아래 RULES에 표기별로 적었고, 애매하면 '검토 필요'로 둔다.
"""
import collections
import csv

import openpyxl

SRC = "data/day1/d1_buyers_raw.xlsx"
DICT = "data/data_dictionary.xlsx"
OUT = "workbench/day1/outputs/d1_code_table.csv"
REVIEW = "검토 필요"

# 키 = 원래 표기를 TRIM · UPPER 한 값. 값 = (표준 코드, 판단 근거, 확신)
RULES = {
    # TT_ADV 선수금
    "T/T IN ADVANCE": ("TT_ADV", "'in advance' = 선적 전 송금", "높음"),
    "100% T/T BEFORE SHIPMENT": ("TT_ADV", "'100%' + 'before shipment' = 전액 선적 전 송금", "높음"),
    "TT": ("TT_ADV", "담당자한테 물어봄 (사람 판단)", "높음"),
    # TT_SPLIT_30_70 선수 + 잔금
    "30% T/T ADVANCE, 70% WITHIN 30 DAYS AFTER B/L": ("TT_SPLIT_30_70", "'30% T/T advance' + '70% within 30 days after B/L' = 코드 정의와 같음", "높음"),
    "T/T 30/70": ("TT_SPLIT_30_70", "'30/70' 분할 비율. 잔금 시점은 표기에 없음(코드의 30일은 교육용 가정)", "높음"),
    "T/T 30:70": ("TT_SPLIT_30_70", "'30:70' 분할 비율. 잔금 시점은 표기에 없음(코드의 30일은 교육용 가정)", "높음"),
    "10% T/T ADVANCE + 90% O/A 90 DAYS": (REVIEW, "'10% advance' + '90% O/A 90 days' — 선수 비율 10/90이 7개 코드 어디에도 없음(TT_SPLIT은 30/70, OA는 선수 0)", "낮음"),
    # OA 사후송금
    "O/A 30": ("OA", "'O/A' + 일수", "높음"),
    "O/A30": ("OA", "'O/A' + 일수(띄어쓰기 없음)", "높음"),
    "OA 30 DAYS": ("OA", "'OA' + 일수", "높음"),
    "30 DAYS OPEN ACCOUNT": ("OA", "'open account' + 일수", "높음"),
    "O/A 60": ("OA", "'O/A' + 일수", "높음"),
    "O/A60": ("OA", "'O/A' + 일수(띄어쓰기 없음)", "높음"),
    "OA 60 DAYS": ("OA", "'OA' + 일수", "높음"),
    "60 DAYS O/A": ("OA", "일수 + 'O/A'", "높음"),
    "OPEN ACCOUNT 60 DAYS": ("OA", "'open account' + 일수", "높음"),
    "O/A 90": ("OA", "'O/A' + 일수", "높음"),
    "OA90": ("OA", "'OA' + 일수(띄어쓰기 없음)", "높음"),
    "OA 90 DAYS": ("OA", "'OA' + 일수", "높음"),
    "OPEN ACCT 90D": ("OA", "'open acct'(open account 약어) + '90d'", "높음"),
    "T/T 30 DAYS AFTER B/L": ("OA", "'T/T' + 'days after B/L' = 선적 후 송금", "높음"),
    "T/T 60 DAYS AFTER B/L": ("OA", "'T/T' + 'days after B/L' = 선적 후 송금", "높음"),
    "T/T 90 DAYS AFTER B/L": ("OA", "'T/T' + 'days after B/L' = 선적 후 송금", "높음"),
    "T/T 60DAYS": ("OA", "'T/T' + 일수 — 사전 codes 시트 OA 설명: \"T/T 60days\"처럼 T/T + 일수 표기는 사후송금", "높음"),
    "NET 30 DAYS AFTER RECEIPT OF INVOICE": ("OA", "'Net 30' + 'after receipt of invoice' = 송장 후 송금(사전 OA 설명의 Net N · 공공 Net30)", "높음"),
    "NET 60": ("OA", "'Net' + 일수(사전 OA 설명의 Net N)", "높음"),
    "NET 90": ("OA", "'Net' + 일수(사전 OA 설명의 Net N)", "높음"),
    # DA 인수인도조건
    "D/A 60": ("DA", "'D/A' + 일수", "높음"),
    "DA60": ("DA", "'DA' + 일수(띄어쓰기 없음)", "높음"),
    "D/A 60 DAYS": ("DA", "'D/A' + 일수", "높음"),
    "D/A 60 DAYS AFTER B/L": ("DA", "'D/A' + 'days after B/L'", "높음"),
    "DOCUMENTS AGAINST ACCEPTANCE 60 DAYS": ("DA", "'documents against acceptance' = D/A", "높음"),
    "D/A 90": ("DA", "'D/A' + 일수", "높음"),
    "DA90": ("DA", "'DA' + 일수(띄어쓰기 없음)", "높음"),
    "D/A 90 DAYS": ("DA", "'D/A' + 일수", "높음"),
    # DP 지급인도조건
    "D/P": ("DP", "'D/P'(일람 여부 표기 없음 — DP 코드는 하나뿐이라 분류에는 영향 없음)", "높음"),
    "D/P AT SIGHT": ("DP", "'D/P' + 'at sight'", "높음"),
    "DP SIGHT": ("DP", "'DP' + 'sight'", "높음"),
    # LC_SIGHT 일람불 신용장
    "L/C AT SIGHT": ("LC_SIGHT", "'L/C' + 'at sight'", "높음"),
    "LC AT SIGHT": ("LC_SIGHT", "'LC' + 'at sight'", "높음"),
    "IRREVOCABLE L/C AT SIGHT": ("LC_SIGHT", "'irrevocable L/C' + 'at sight'", "높음"),
    "CONFIRMED L/C AT SIGHT": ("LC_SIGHT", "'confirmed L/C' + 'at sight'(확인 여부는 코드에 반영 안 됨)", "높음"),
    "신용장 일람불": ("LC_SIGHT", "'신용장' + '일람불' = sight L/C", "높음"),
    # LC_USANCE 기한부 신용장
    "USANCE L/C 90 DAYS": ("LC_USANCE", "'usance L/C' + 일수", "높음"),
    "L/C 60 DAYS": ("LC_USANCE", "'L/C' + 일수. 'usance'라는 말은 없음 — 신용장 유효기간일 가능성도 있음", "낮음"),
    "L/C 90 DAYS": ("LC_USANCE", "'L/C' + 일수. 'usance'라는 말은 없음 — 신용장 유효기간일 가능성도 있음", "낮음"),
}

CODE_ORDER = ["TT_ADV", "TT_SPLIT_30_70", "LC_SIGHT", "LC_USANCE", "DP", "DA", "OA", REVIEW]

wb = openpyxl.load_workbook(SRC, read_only=True)
rows = list(wb["buyers"].iter_rows(values_only=True))
j = rows[0].index("payment_terms_raw")
raw_count = collections.Counter(r[j] for r in rows[1:])

missing = sorted({r.strip().upper() for r in raw_count} - RULES.keys())
if missing:
    raise SystemExit(f"RULES에 없는 표기: {missing}")

table = []
for raw, n in raw_count.items():
    code, why, conf = RULES[raw.strip().upper()]
    if raw != raw.strip():
        why += " · 원본에 앞뒤 공백 있음"
    table.append((raw, code, why, conf, n))
table.sort(key=lambda t: (CODE_ORDER.index(t[1]), t[0].strip().upper(), t[0]))

with open(OUT, "w", encoding="utf-8-sig", newline="") as f:
    w = csv.writer(f)
    w.writerow(["원래 표기", "표준 코드", "판단 근거", "확신"])
    w.writerows(t[:4] for t in table)

print(f"표기 {len(table)}가지(TRIM · UPPER 후 {len({t[0].strip().upper() for t in table})}가지), 셀 {sum(raw_count.values())}개 → {OUT}")
by_code = collections.Counter()
for t in table:
    by_code[t[1]] += t[4]
print("코드별 셀 수:", {c: by_code[c] for c in CODE_ORDER})
print("확신 낮음:", [(t[0], t[1], t[4]) for t in table if t[3] == "낮음"])

# 사전의 참고 코드표(payment_terms_map)와 대조 — 다르거나 참고표에 없는 것만 출력
ref = {r[1]: r[2] for r in openpyxl.load_workbook(DICT, read_only=True)["payment_terms_map"]
       .iter_rows(min_row=4, values_only=True) if r[1]}
for key, (code, _, _) in sorted(RULES.items()):
    if key not in ref:
        print(f"참고표에 없음: {key!r} → 내 판단 {code}")
    elif ref[key] != code:
        print(f"참고표와 다름: {key!r} → 내 판단 {code} / 참고표 {ref[key]}")
